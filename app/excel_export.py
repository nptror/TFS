"""Tạo .xlsx trong bộ nhớ, không lưu dữ liệu học sinh lên ổ đĩa."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from .formatter import format_title
from .validator import absence_count, class_sort_key


HEADERS = ["Lớp", "Sĩ số", "Số vắng", "Có phép", "Không phép", "Danh sách vắng", "Ghi chú"]


def _set_text(cell, value):
    """Giữ nội dung là text, kể cả khi người dùng nhập ký tự = đầu chuỗi."""
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell.data_type = "s"


def export_excel(data: dict, warnings: list[dict], buoi: str, ngay: date, ca: int) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Báo vắng"
    sheet.merge_cells("A1:G1")
    _set_text(sheet["A1"], format_title(buoi, ngay, ca))
    sheet["A1"].font = Font(name="Calibri", size=15, bold=True, color="164E42")
    sheet["A1"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34
    sheet.append(HEADERS)
    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245C4F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 26
    last_grade = None
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for record in sorted(data["lop"], key=class_sort_key):
        grade = record["ten_lop"][:2]
        if last_grade is not None and grade != last_grade:
            sheet.append([None] * 7)
        students = record["hoc_sinh"]
        notes = [w["message"] + ("\n" + w["raw"] if w.get("raw") else "")
                 for w in warnings if w.get("class_name") == record["ten_lop"]]
        values = [record["ten_lop"], record["si_so"], absence_count(record),
                  sum(s["loai"] == "P" for s in students) if absence_count(record) is not None else None,
                  sum(s["loai"] == "K" for s in students) if absence_count(record) is not None else None,
                  ", ".join(f"{s['ten']} ({s['loai'] or '?'})" for s in students), "\n".join(notes)]
        sheet.append(values)
        for column, value in enumerate(values, 1):
            cell = sheet.cell(sheet.max_row, column)
            if isinstance(value, str):
                _set_text(cell, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if notes:
                cell.fill = yellow
        sheet.row_dimensions[sheet.max_row].height = max(30, 15 * (1 + len(notes) + len(values[5]) // 65))
        last_grade = grade
    # Cảnh báo ngoài lớp không bị bỏ mất khi chuyển sang Excel.
    global_notes = [w for w in warnings if not w.get("class_name")]
    if global_notes:
        sheet.append([None] * 7)
        for warning in global_notes:
            sheet.append(["Cần kiểm tra", None, None, None, None, None, None])
            cell = sheet.cell(sheet.max_row, 7)
            _set_text(cell, warning["message"] + "\n" + (warning.get("raw") or ""))
            cell.fill = yellow
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[sheet.max_row].height = 45
    for index, width in enumerate([14, 12, 12, 12, 14, 72, 64], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "C3"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    result = BytesIO()
    workbook.save(result)
    result.seek(0)
    return result


def export_filename(buoi: str, ngay: date, ca: int) -> str:
    return f"bao-vang_{ngay.isoformat()}_{'sang' if buoi == 'Sáng' else 'chieu'}_ca{ca}.xlsx"
