"""Nhập điểm từ phiếu scan: PDF + Excel trống -> duyệt -> Excel có điểm.

Lõi OCR (tách bảng, Tesseract, mô hình chữ số) nằm trong app/ocr.
"""

import json
import os
import subprocess
import tempfile
import zipfile
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from typing import Annotated
from urllib.parse import quote

import openpyxl
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from starlette.concurrency import run_in_threadpool


ROOT = Path(__file__).resolve().parent.parent
MAX_EXCEL_BYTES = 10 * 1024 * 1024
MAX_PDF_BYTES = 40 * 1024 * 1024
MAX_PAGES = 30
THRESHOLD = 0.85
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

router = APIRouter()


def _core():
    # Nạp muộn: OpenCV/Tesseract nặng và tuỳ chọn; trang báo vắng vẫn chạy khi chưa cài.
    try:
        from .ocr import excel, grid, pipeline
    except ImportError as error:
        raise HTTPException(503, "Chưa cài thư viện nhập điểm. Chạy: pip install -r requirements-nhapdiem.txt") from error
    return excel, grid, pipeline


@lru_cache(maxsize=1)
def _model():
    return _core()[2].load_model()


async def _read_upload(upload: UploadFile, limit: int, label: str) -> bytes:
    try:
        data = await upload.read(limit + 1)
    finally:
        await upload.close()
    if not data:
        raise HTTPException(400, f"File {label} trống.")
    if len(data) > limit:
        raise HTTPException(400, f"File {label} vượt quá {limit // (1024 * 1024)} MB.")
    return data


def _open_workbook(data: bytes, filename: str, folder: str) -> str:
    excel = _core()[0]
    path = os.path.join(folder, "dau_vao.xls" if filename.lower().endswith(".xls") else "dau_vao.xlsx")
    Path(path).write_bytes(data)
    try:
        path = excel.ensure_xlsx(path)
        openpyxl.load_workbook(path, read_only=True).close()
    except RuntimeError as error:
        raise HTTPException(400, str(error)) from error
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, subprocess.SubprocessError) as error:
        raise HTTPException(400, "Không mở được file Excel. Hãy dùng file .xlsx hợp lệ.") from error
    return path


def _sheet(path: str, sheet: str, subject: str):
    excel = _core()[0]
    workbook = openpyxl.load_workbook(path)
    if sheet not in workbook.sheetnames:
        raise HTTPException(400, f"Không có sheet '{sheet}' trong file Excel.")
    try:
        students, _, _, subjects = excel.load_students(workbook[sheet])
    except ValueError as error:
        raise HTTPException(400, str(error)) from error
    if subject not in subjects:
        raise HTTPException(400, f"Sheet '{sheet}' không có cột môn '{subject}'.")
    return workbook[sheet], students, subjects[subject]


def _list_sheets(path: str) -> list[dict]:
    excel = _core()[0]
    sheets = []
    for worksheet in openpyxl.load_workbook(path).worksheets:
        try:
            students, _, _, subjects = excel.load_students(worksheet)
        except (ValueError, KeyError):
            continue
        if subjects:
            sheets.append({"name": worksheet.title, "subjects": list(subjects), "students": len(students)})
    if not sheets:
        raise HTTPException(400, "Không tìm thấy sheet có dòng tiêu đề 'Họ và tên' và cột môn sau cột 'Phòng'.")
    return sheets


def _flags(row: dict, existing: dict) -> list[str]:
    flags = []
    if row["trong"]:
        flags.append("vắng/trống")
    elif row["diem"] is None:
        flags.append("không đọc được")
    elif row["conf"] < THRESHOLD:
        flags.append("điểm chưa chắc")
    if row["match_conf"] < 0.8 or row["match_note"]:
        flags.append(row["match_note"] or "tên chưa chắc")
    student = row["hs"]
    if student and existing.get(student["row"]) not in (None, ""):
        flags.append(f"đã có điểm {existing[student['row']]}")
    return flags


def _read_scores(path: str, pdf: bytes, sheet: str, subject: str) -> dict:
    import pytesseract

    _, grid, pipeline = _core()
    worksheet, students, column = _sheet(path, sheet, subject)
    try:
        images = grid.pdf_to_images(pdf)
    except Exception as error:  # PyMuPDF ném nhiều loại lỗi khác nhau cho PDF hỏng
        raise HTTPException(400, "Không đọc được file PDF. Hãy kiểm tra lại file scan.") from error
    if not images:
        raise HTTPException(400, "File PDF không có trang nào.")
    if len(images) > MAX_PAGES:
        raise HTTPException(400, f"Chỉ nhận tối đa {MAX_PAGES} trang mỗi lần.")
    try:
        rows, errors = pipeline.process_pages(images, students, _model())
    except pytesseract.TesseractNotFoundError as error:
        raise HTTPException(503, "Chưa cài Tesseract OCR trên máy chủ. Xem hướng dẫn trong README.") from error
    existing = {s["row"]: worksheet.cell(s["row"], column).value for s in students}
    return {
        "students": [{"row": s["row"], "label": f"{s['ho_ten']} · {s['ngay_sinh']} · {s['lop']} (dòng {s['row']})",
                      "ho_ten": s["ho_ten"], "ngay_sinh": s["ngay_sinh"], "lop": s["lop"],
                      "existing": None if existing[s["row"]] in (None, "") else str(existing[s["row"]])}
                     for s in students],
        "rows": [{
            "trang": row["trang"], "stt": row["stt"],
            "anh_ten": row["anh_ten"], "anh_diem": row["anh_diem"],
            "diem": None if row["diem"] is None else float(row["diem"]),
            "conf": float(row["conf"]),
            "row": row["hs"]["row"] if row["hs"] else None,
            "flags": _flags(row, existing),
        } for row in rows],
        "errors": [{"trang": page, "message": message} for page, message in errors],
    }


def _parse_items(items: str, students: list[dict]) -> list[tuple[int, float]]:
    try:
        parsed = [(int(row), round(float(score), 2)) for row, score in json.loads(items)]
    except (ValueError, TypeError) as error:
        raise HTTPException(400, "Danh sách điểm không hợp lệ.") from error
    valid_rows = {s["row"] for s in students}
    rows = [row for row, _ in parsed]
    if any(row not in valid_rows for row in rows):
        raise HTTPException(400, "Có dòng Excel không thuộc danh sách học sinh.")
    if len(set(rows)) != len(rows):
        raise HTTPException(400, "Có học sinh bị ghép 2 lần; hãy sửa trước khi xuất.")
    if any(not 0 <= score <= 10 for _, score in parsed):
        raise HTTPException(400, "Điểm phải nằm trong khoảng 0–10.")
    return parsed


def _export(path: str, sheet: str, subject: str, items: str) -> bytes:
    excel = _core()[0]
    _, students, column = _sheet(path, sheet, subject)
    output = os.path.join(os.path.dirname(path), "da_nhap.xlsx")
    excel.write_scores(path, output, sheet, column, _parse_items(items, students))
    return Path(output).read_bytes()


@router.get("/nhapdiem", include_in_schema=False)
async def page():
    return FileResponse(ROOT / "app" / "static" / "nhapdiem.html")


@router.post("/api/nhapdiem/sheets")
async def sheets(excel: Annotated[UploadFile, File()]):
    data = await _read_upload(excel, MAX_EXCEL_BYTES, "Excel")
    with tempfile.TemporaryDirectory() as folder:
        path = await run_in_threadpool(_open_workbook, data, excel.filename or "", folder)
        return {"sheets": await run_in_threadpool(_list_sheets, path)}


@router.post("/api/nhapdiem/read")
async def read(excel: Annotated[UploadFile, File()], pdf: Annotated[UploadFile, File()],
               sheet: Annotated[str, Form(max_length=200)], subject: Annotated[str, Form(max_length=200)]):
    data = await _read_upload(excel, MAX_EXCEL_BYTES, "Excel")
    pdf_data = await _read_upload(pdf, MAX_PDF_BYTES, "PDF")
    with tempfile.TemporaryDirectory() as folder:
        path = await run_in_threadpool(_open_workbook, data, excel.filename or "", folder)
        return await run_in_threadpool(_read_scores, path, pdf_data, sheet, subject)


@router.post("/api/nhapdiem/export")
async def export(excel: Annotated[UploadFile, File()], sheet: Annotated[str, Form(max_length=200)],
                 subject: Annotated[str, Form(max_length=200)], items: Annotated[str, Form(max_length=200_000)]):
    filename = excel.filename or "diem.xlsx"
    data = await _read_upload(excel, MAX_EXCEL_BYTES, "Excel")
    with tempfile.TemporaryDirectory() as folder:
        path = await run_in_threadpool(_open_workbook, data, filename, folder)
        content = await run_in_threadpool(_export, path, sheet, subject, items)
    name = f"{Path(filename).stem}_{subject}_da_nhap.xlsx"
    return StreamingResponse(BytesIO(content), media_type=XLSX, headers={
        "Content-Disposition": f"attachment; filename=\"diem_da_nhap.xlsx\"; filename*=UTF-8''{quote(name)}"})
