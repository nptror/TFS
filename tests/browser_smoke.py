"""Optional UI smoke test: start uvicorn on 8765, then run this file.

Requires playwright and a locally installed Chrome. No external API is called.
"""

from io import BytesIO
from pathlib import Path

from PIL import Image
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / ".tools" / "browser"


def run():
    OUT.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(channel="chrome", headless=True)
        context = browser.new_context(viewport={"width": 1440, "height": 1000}, locale="vi-VN",
                                      permissions=["clipboard-read", "clipboard-write"])
        page = context.new_page()
        errors = []
        page.on("pageerror", lambda error: errors.append(str(error)))
        page.goto("http://127.0.0.1:8765", wait_until="networkidle")
        expect(page.locator("#config-count")).to_have_text("(14 lớp)")
        expect(page.locator("#upload-section")).to_be_hidden()
        page.locator("#session").select_option("Chiều")
        page.locator("#date").fill("2026-10-23")
        page.get_by_role("button", name="Dùng ví dụ").click()
        page.locator("#input-text").press("Control+Enter")
        expect(page.locator("#result-state")).to_have_text("Đã chuẩn hóa")
        expect(page.locator("#stat-classes")).to_have_text("7/14")
        expect(page.locator("#stat-absent")).to_have_text("19")
        assert "Huy Hoàng (P)" in page.locator("#output-text").input_value()
        assert "Chưa rõ P/K: Phương Thảo" in page.locator("#warning-list").inner_text()
        page.screenshot(path=str(OUT / "desktop.png"), full_page=True)
        page.get_by_role("button", name="Copy", exact=True).click()
        expect(page.locator("#message")).to_have_text("Đã copy bản báo cáo.")
        copied = page.evaluate("navigator.clipboard.readText()")
        assert copied.replace("\r\n", "\n") == page.locator("#output-text").input_value()
        original = page.locator("#output-text").input_value()
        page.locator("#output-text").fill(original.replace("Phương Thảo (?)", "Phương Thảo (P)"))
        expect(page.locator("#edit-status")).to_have_text("Đã kiểm tra lại bản chỉnh sửa")
        assert "Chưa rõ P/K: Phương Thảo" not in page.locator("#warning-list").inner_text()
        with page.expect_download() as download_info:
            page.get_by_role("button", name="Xuất Excel", exact=True).click()
        download = download_info.value
        assert download.suggested_filename == "bao-vang_2026-10-23_chieu_ca1.xlsx"
        download.save_as(str(OUT / download.suggested_filename))
        sheet = load_workbook(OUT / download.suggested_filename).active
        row = next(row for row in sheet.iter_rows() if row[0].value == "12C11")
        assert row[3].value == 1 and row[5].value == "Phương Thảo (P)"
        row = next(row for row in sheet.iter_rows() if row[0].value == "12C2")
        assert row[1].value == 40
        expect(page.locator("#export-button")).to_be_enabled()
        page.reload(wait_until="networkidle")
        expect(page.locator("#edit-status")).to_have_text("Đã kiểm tra lại bản chỉnh sửa")
        assert "Phương Thảo (P)" in page.locator("#output-text").input_value()
        assert page.locator("#date").input_value() == "2026-10-23"

        # Upload and remove thumbnails without sending any data to Gemini.
        page.get_by_text("Có ảnh", exact=True).click()
        expect(page.locator("#upload-section")).to_be_visible()
        image = BytesIO()
        Image.new("RGB", (20,20), "white").save(image, format="PNG")
        page.locator("#image-input").set_input_files([
            {"name":"first.png", "mimeType":"image/png", "buffer":image.getvalue()},
            {"name":"second.png", "mimeType":"image/png", "buffer":image.getvalue()},
        ])
        expect(page.locator(".thumbnail")).to_have_count(2)
        page.get_by_role("button", name="Xóa ảnh first.png").click()
        expect(page.locator(".thumbnail")).to_have_count(1)
        page.get_by_text("Chỉ text", exact=True).click()
        expect(page.locator("#upload-section")).to_be_hidden()

        for width in [390, 320, 768, 1024]:
            page.set_viewport_size({"width": width, "height":844})
            assert page.evaluate("document.documentElement.scrollWidth <= innerWidth"), f"Overflow at {width}px"
            if width < 800:
                left = page.locator("#input-title").bounding_box()
                control = page.locator("#control-title").bounding_box()
                right = page.locator("#output-title").bounding_box()
                assert left["y"] < control["y"] < right["y"]
            if width == 390:
                page.screenshot(path=str(OUT / "mobile.png"), full_page=True)
        page.get_by_role("button", name="Xóa nội dung").click()
        expect(page.locator("#input-text")).to_have_value("")
        expect(page.locator("#output-text")).to_have_value("")
        expect(page.locator("#export-button")).to_be_disabled()
        page.reload(wait_until="networkidle")
        expect(page.locator("#input-text")).to_have_value("")
        assert not errors, errors
        context.close()
        browser.close()
    print("Browser smoke passed: format, keyboard, editing, copy, Excel, images, draft, responsive.")


if __name__ == "__main__":
    run()
