from io import BytesIO
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest

from app import main


SESSION = {"buoi": "Chiều", "ngay": "2026-10-23", "ca": 1}


def format_request(client, text):
    return client.post("/api/format", data={**SESSION, "text": text})


def export_payload(result, text=None):
    return {**SESSION, "text": text if text is not None else result["text"], "source_data": result["data"]}


def class_row(sheet, code):
    return next(row for row in sheet.iter_rows() if row[0].value == code)


def test_home_and_config(client):
    assert client.get("/").status_code == 200
    assert "Tin nhắn đầu vào" in client.get("/").text
    assert "10C10" in client.get("/api/config").json()["classes"]
    response = format_request(client, "12C9 vắng 1 An (P)")
    assert response.status_code == 200
    assert "12C9 Vắng  1 An (P)" in response.json()["text"]


def test_export_manual_edits_metadata_and_original_size(client):
    response = format_request(client, "12C9 sĩ số 40 vắng 2 An(P), Bình(K)")
    result = response.json()
    edited = result["text"].replace("2 An (P), Bình (K)", "1 Huy (K)").replace(
        "Chiều thứ 6, ngày 23/10/2026, ca 1:", "Sáng Chủ nhật, ngày 25/10/2026, ca 3:")
    exported = client.post("/api/export", json=export_payload(result, edited))
    assert exported.status_code == 200
    assert 'bao-vang_2026-10-25_sang_ca3.xlsx' in exported.headers["content-disposition"]
    sheet = load_workbook(BytesIO(exported.content)).active
    assert sheet["A1"].value == "Sáng Chủ nhật, ngày 25/10/2026, ca 3"
    assert sheet["A1"].font.bold
    assert "A1:G1" in str(sheet.merged_cells)
    row = class_row(sheet, "12C9")
    assert [c.value for c in row] == ["12C9", 40, 1, 0, 1, "Huy (K)", None]
    assert sheet.column_dimensions["F"].width >= 60
    assert any(all(c.value is None for c in row) for row in sheet.iter_rows(min_row=3))
    missing = class_row(sheet, "10C10")
    assert missing[2].value is None and "Chưa báo" in missing[6].value
    assert missing[6].fill.fgColor.rgb == "00FFF2CC"


def test_export_preserves_source_discrepancy_until_line_edited(client):
    result = format_request(client, "12C9 vắng 3 An(P), Bình(K)").json()
    unchanged = client.post("/api/preview", json=export_payload(result)).json()
    assert any(w["code"] == "count_mismatch" and w["class_name"] == "12C9" for w in unchanged["warnings"])
    edited = result["text"].replace("2 An (P), Bình (K)", "1 An (P)")
    changed = client.post("/api/preview", json=export_payload(result, edited)).json()
    assert not any(w["code"] == "count_mismatch" for w in changed["warnings"])
    # Sửa riêng tổng cũng phải được nhận diện và kiểm tra, không lấy lại tổng cũ.
    wrong_total = result["text"].replace("2 An (P)", "99 An (P)")
    changed = client.post("/api/preview", json=export_payload(result, wrong_total)).json()
    assert any("99 vắng" in w["message"] for w in changed["warnings"])


def test_export_retains_unclear_and_duplicate_warnings(client):
    result = format_request(client, "chào\n12C9 vắng 0\n12C9 vắng 1 an").json()
    response = client.post("/api/export", json=export_payload(result))
    sheet = load_workbook(BytesIO(response.content)).active
    row = class_row(sheet, "12C9")
    assert "Đã ghi đè" in row[6].value
    assert row[5].value == "An (K)"
    assert any(row[6].value and "chào" in row[6].value for row in sheet.iter_rows())


@pytest.mark.parametrize("updates", [{"text": ""}, {"ca": 0}, {"ngay": "2026-02-31"}])
def test_invalid_requests(client, updates):
    data = {**SESSION, "text": "12C9 vắng 0", **updates}
    assert client.post("/api/format", data=data).status_code in (400, 422)


def test_bad_config_reports_clear_error(client, monkeypatch, tmp_path):
    path = tmp_path / "config.json"
    path.write_text(json.dumps({"classes": ["typo"]}), encoding="utf-8")
    monkeypatch.setattr(main, "CONFIG_PATH", path)
    response = client.get("/api/config")
    assert response.status_code == 500 and "config.json" in response.json()["detail"]


def test_export_invalid_header_and_missing_classes(client):
    assert client.post("/api/export", json={**SESSION, "text": "không có lớp"}).status_code == 400
    assert client.post("/api/export", json={**SESSION, "text": "Sáng thứ 2, ngày 31/02/2026, ca 1:\n12C9 vắng 0"}).status_code == 422


def test_extended_history_through_format_preview_and_excel(client):
    text = (Path(__file__).parent / "fixtures" / "chat_history.txt").read_text(encoding="utf-8")
    response = format_request(client, text)
    assert response.status_code == 200
    result = response.json()
    assert result["summary"]["reported"] == 25
    assert "12C6 Vắng  0" in result["text"]
    assert "11C1 Vắng  2 Tấn Phát (K), Nhật Uy (K)" in result["text"]
    assert "Quang Vinh" not in result["text"]
    payload = export_payload(result)
    preview = client.post("/api/preview", json=payload)
    assert preview.status_code == 200
    assert preview.json()["text"] == result["text"]
    response = client.post("/api/export", json=payload)
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    row = class_row(sheet, "12C14")
    assert row[2].value == 0
    assert "Quang Vinh" not in (row[6].value or "")
    assert "đi trễ" not in (row[6].value or "")
    assert class_row(sheet, "12C10")[1].value == 43


def test_all_late_students_export_as_zero_but_keep_discrepancy(client):
    result = format_request(client, "12C1 Vắng 2 An(T), Bình(T 1t)").json()
    assert "12C1 Vắng  0" in result["text"]
    preview = client.post("/api/preview", json=export_payload(result))
    assert preview.status_code == 200
    assert "12C1 Vắng  0" in preview.json()["text"]
    codes = {w["code"] for w in preview.json()["warnings"] if w["class_name"] == "12C1"}
    assert "count_mismatch" in codes
    assert "late_excluded" not in codes
    assert all("Bình" not in warning["message"] for warning in preview.json()["warnings"])


def test_group_without_count_formats_and_exports(client):
    result = format_request(client, "12C18 Vắng không phép: An, Bình").json()
    assert "12C18 Vắng  2 An (K), Bình (K)" in result["text"]
    assert client.post("/api/export", json=export_payload(result)).status_code == 200

