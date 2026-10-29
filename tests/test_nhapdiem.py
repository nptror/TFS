import json
from io import BytesIO

import openpyxl
import pytest

pytest.importorskip("cv2")  # thư viện OCR là tuỳ chọn (requirements-nhapdiem.txt)

from app import nhapdiem  # noqa: E402


def blank_workbook() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "TỔNG 11"
    sheet.append(["Danh sách thi cuối kỳ"])
    sheet.append(["STT", "SBD", "Họ và tên", "Ngày sinh", "Lớp", "Phòng", "KTPL", "Toán"])
    sheet.append([1, "001", "Nguyễn Văn An", "01/02/2010", "11C1", "P1", None, 7])
    sheet.append([2, "002", "Trần Thị Bình", "03/04/2010", "11C1", "P1", None, None])
    workbook.create_sheet("Ghi chú").append(["không phải danh sách"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def excel_file():
    return ("excel", ("ds.xlsx", blank_workbook(), nhapdiem.XLSX))


def test_page_and_sheets(client):
    assert client.get("/nhapdiem").status_code == 200
    response = client.post("/api/nhapdiem/sheets", files=[excel_file()])
    assert response.status_code == 200
    assert response.json()["sheets"] == [{"name": "TỔNG 11", "subjects": ["KTPL", "TOÁN"], "students": 2}]
    bad = client.post("/api/nhapdiem/sheets", files=[("excel", ("x.xlsx", b"not excel", nhapdiem.XLSX))])
    assert bad.status_code == 400


def test_read_returns_rows_with_flags(client, monkeypatch):
    _, grid, pipeline = nhapdiem._core()
    monkeypatch.setattr(grid, "pdf_to_images", lambda data: ["page"])
    monkeypatch.setattr(nhapdiem, "_model", lambda: None)

    def fake_process(images, students, model):
        row = {"trang": 1, "stt": 1, "anh_ten": "data:,", "anh_diem": "data:,", "trong": False,
               "match_note": "", "_score_cell": object()}
        return [{**row, "diem": 8.5, "conf": 0.95, "hs": students[0], "match_conf": 0.9},
                {**row, "stt": 2, "diem": 6.0, "conf": 0.4, "hs": students[1], "match_conf": 0.9}], [(2, "không có bảng")]
    monkeypatch.setattr(pipeline, "process_pages", fake_process)
    data = {"sheet": "TỔNG 11", "subject": "TOÁN"}
    files = [excel_file(), ("pdf", ("p.pdf", b"%PDF", "application/pdf"))]
    result = client.post("/api/nhapdiem/read", data=data, files=files).json()
    assert [(r["row"], r["diem"], r["flags"]) for r in result["rows"]] == [
        (3, 8.5, ["đã có điểm 7"]), (4, 6.0, ["điểm chưa chắc"])]
    assert result["errors"] == [{"trang": 2, "message": "không có bảng"}]
    assert [(s["row"], s["ho_ten"], s["existing"]) for s in result["students"]] == [
        (3, "Nguyễn Văn An", "7"), (4, "Trần Thị Bình", None)]
    assert client.post("/api/nhapdiem/read", data={**data, "subject": "Văn"}, files=files).status_code == 400


def test_export_writes_scores_and_validates(client):
    data = {"sheet": "TỔNG 11", "subject": "KTPL"}
    response = client.post("/api/nhapdiem/export", data={**data, "items": json.dumps([[3, 8.5], [4, 10]])},
                           files=[excel_file()])
    assert response.status_code == 200
    assert "ds_KTPL_da_nhap.xlsx" in response.headers["content-disposition"]
    sheet = openpyxl.load_workbook(BytesIO(response.content))["TỔNG 11"]
    assert [sheet.cell(r, 7).value for r in (3, 4)] == [8.5, 10]
    assert sheet.cell(3, 8).value == 7  # cột môn khác giữ nguyên
    for items in ([[3, 8], [3, 9]], [[3, 11]], [[99, 5]], "rác"):
        payload = {**data, "items": items if isinstance(items, str) else json.dumps(items)}
        assert client.post("/api/nhapdiem/export", data=payload, files=[excel_file()]).status_code == 400
